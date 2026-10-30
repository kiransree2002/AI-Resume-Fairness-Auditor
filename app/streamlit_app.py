
import io
import sys
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from uuid import uuid4

import pandas as pd
import plotly.graph_objects as go
import streamlit as st
import pdfplumber
from docx import Document

# Make the src folder importable when Streamlit runs this file.
PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from jd_parser import parse_job_description, scrape_job_url
from resume_parser import parse_resume
from gate_cascade import run_candidate_cascade
from semantic_scorer import (
    SemanticScorer,
)


st.set_page_config(
    page_title="TalentLens | Recruitment intelligence",
    page_icon=":material/description:",
    layout="wide",
)

st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700;800&family=Space+Grotesk:wght@500;600;700&display=swap');

    :root {
        --audit-navy: #0B1120;
        --audit-navy-soft: #11182E;
        --audit-blue: #0284C7;
        --audit-indigo: #4F46E5;
        --audit-teal: #0D9488;
        --audit-text: #1E293B;
        --audit-muted: #64748B;
        --audit-border: #E2E8F0;
        --audit-surface: #FFFFFF;
    }

    .stApp {
        background:
            radial-gradient(ellipse at 8% 0%, rgba(99, 102, 241, 0.075), transparent 34%),
            radial-gradient(ellipse at 96% 12%, rgba(20, 184, 166, 0.065), transparent 30%),
            #F8FAFC;
        color: var(--audit-text);
        font-family: 'Plus Jakarta Sans', -apple-system, BlinkMacSystemFont, sans-serif;
    }

    .block-container {
        max-width: 1560px;
        padding-top: 1.8rem;
        padding-bottom: 3.2rem;
    }

    h1, h2, h3, h4,
    [data-testid="stMetricValue"] {
        color: #0F172A;
        font-family: 'Space Grotesk', 'Plus Jakarta Sans', sans-serif;
        font-weight: 700;
    }

    h1 { letter-spacing: -0.035em; }
    h2, h3 { letter-spacing: -0.02em; }

    [data-testid="stSidebar"] {
        background: linear-gradient(175deg, #0B1020 0%, #11182E 52%, #152238 100%);
        border-right: 1px solid rgba(255, 255, 255, 0.08);
    }

    [data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p,
    [data-testid="stSidebar"] [data-testid="stRadio"] label {
        color: #EEF2FF;
    }

    [data-testid="stSidebar"] [data-testid="stCaptionContainer"] {
        color: #94A3B8;
    }

    [data-testid="stSidebar"] [data-testid="stRadio"] label {
        border: 1px solid transparent;
        border-radius: 12px;
        padding: 0.42rem 0.65rem;
        font-weight: 600;
        transition: background-color 180ms ease, border-color 180ms ease, transform 180ms ease;
    }

    [data-testid="stSidebar"] [data-testid="stRadio"] label:hover {
        background: rgba(99, 102, 241, 0.18);
        border-color: rgba(129, 140, 248, 0.25);
        transform: translateX(2px);
    }

    .hero-card {
        position: relative;
        overflow: hidden;
        padding: 1.5rem 1.8rem;
        margin-bottom: 1.4rem;
        border: 1px solid rgba(255, 255, 255, 0.24);
        border-radius: 20px;
        color: #FFFFFF;
        background: linear-gradient(120deg, #1E1B4B 0%, #312E81 35%, #4F46E5 72%, #0D9488 125%);
        box-shadow: 0 16px 36px rgba(79, 70, 229, 0.18);
        animation: workspace-enter 240ms ease-out both;
    }

    .hero-title {
        color: #FFFFFF;
        font-family: 'Space Grotesk', sans-serif;
        font-size: clamp(1.35rem, 2.4vw, 1.8rem);
        font-weight: 700;
        letter-spacing: -0.025em;
        margin-bottom: 0.25rem;
    }

    .hero-subtitle {
        color: rgba(248, 250, 252, 0.88);
        font-size: 0.9rem;
    }

    .st-key-overview-hero {
        padding: clamp(1.7rem, 4vw, 3.2rem);
        border: 1px solid rgba(226, 232, 240, 0.9);
        border-radius: 24px;
        background: linear-gradient(145deg, rgba(255,255,255,0.98), rgba(241,245,249,0.96));
        box-shadow: 0 18px 48px rgba(15, 23, 42, 0.055);
        animation: workspace-enter 240ms ease-out both;
    }

    [data-testid="stMetric"] {
        position: relative;
        border: 1px solid var(--audit-border);
        border-radius: 16px;
        padding: 1rem 1.1rem;
        background: linear-gradient(155deg, #FFFFFF, #F8FAFC);
        box-shadow: 0 4px 16px rgba(15, 23, 42, 0.04);
        transition: transform 200ms cubic-bezier(0.16, 1, 0.3, 1),
                    border-color 200ms ease, box-shadow 200ms ease;
    }

    [data-testid="stMetric"]:hover {
        transform: translateY(-2px);
        border-color: rgba(99, 102, 241, 0.35);
        box-shadow: 0 12px 26px rgba(79, 70, 229, 0.11);
    }

    [data-testid="stMetricLabel"] {
        color: var(--audit-muted);
        font-size: 0.85rem;
        font-weight: 600;
    }

    [data-testid="stVerticalBlockBorderWrapper"] {
        border-color: var(--audit-border);
        border-radius: 16px;
        background: rgba(255, 255, 255, 0.97);
        box-shadow: 0 6px 20px rgba(15, 23, 42, 0.035);
        transition: transform 200ms cubic-bezier(0.16, 1, 0.3, 1),
                    border-color 200ms ease, box-shadow 200ms ease;
    }

    [data-testid="stVerticalBlockBorderWrapper"]:hover {
        transform: translateY(-2px);
        border-color: rgba(13, 148, 136, 0.32);
        box-shadow: 0 12px 28px rgba(15, 23, 42, 0.07);
    }

    .stButton > button,
    .stDownloadButton > button {
        min-height: 2.65rem;
        border: 1px solid #4F46E5;
        border-radius: 12px;
        background: linear-gradient(135deg, #4F46E5 0%, #6366F1 58%, #0D9488 140%);
        color: #FFFFFF;
        font-family: 'Plus Jakarta Sans', sans-serif;
        font-weight: 700;
        box-shadow: 0 4px 14px rgba(79, 70, 229, 0.18);
        transition: transform 180ms cubic-bezier(0.16, 1, 0.3, 1),
                    box-shadow 180ms ease, filter 180ms ease;
    }

    .stButton > button:hover,
    .stDownloadButton > button:hover {
        transform: translateY(-2px);
        filter: brightness(1.04);
        box-shadow: 0 8px 22px rgba(79, 70, 229, 0.27);
    }

    .stButton > button:focus-visible,
    .stDownloadButton > button:focus-visible,
    a:focus-visible {
        outline: 3px solid rgba(13, 148, 136, 0.5);
        outline-offset: 3px;
    }

    .stTextInput input,
    .stTextArea textarea,
    .stNumberInput input,
    [data-baseweb="select"] > div {
        border-color: #CBD5E1;
        border-radius: 11px;
        background: #FFFFFF;
        transition: border-color 160ms ease, box-shadow 160ms ease;
    }

    .stTextInput input:focus,
    .stTextArea textarea:focus,
    .stNumberInput input:focus {
        border-color: #6366F1;
        box-shadow: 0 0 0 3px rgba(99, 102, 241, 0.14);
    }

    [data-testid="stFileUploaderDropzone"] {
        border: 1.5px dashed #A5B4C7;
        border-radius: 14px;
        background: linear-gradient(135deg, #FFFFFF, #F8FAFC);
        transition: border-color 180ms ease, background-color 180ms ease;
    }

    [data-testid="stFileUploaderDropzone"]:hover {
        border-color: #0D9488;
        background: #F0FDFA;
    }

    [data-testid="stDataFrame"],
    [data-testid="stDataEditor"] {
        overflow: hidden;
        border: 1px solid var(--audit-border);
        border-radius: 14px;
        box-shadow: 0 4px 16px rgba(15, 23, 42, 0.04);
    }

    [data-testid="stExpander"] {
        overflow: hidden;
        border: 1px solid var(--audit-border);
        border-radius: 14px;
        background: #FFFFFF;
        box-shadow: 0 2px 8px rgba(15, 23, 42, 0.025);
    }

    [data-testid="stExpander"] summary {
        color: #1E293B;
        font-weight: 700;
    }

    [data-testid="stTabs"] [data-baseweb="tab-list"] {
        gap: 0.45rem;
        border-bottom: 1px solid #DCE3ED;
    }

    [data-testid="stTabs"] [data-baseweb="tab"] {
        border-radius: 10px 10px 0 0;
        font-weight: 700;
        transition: color 160ms ease, background-color 160ms ease;
    }

    [data-testid="stTabs"] [aria-selected="true"] {
        color: #4338CA;
        border-bottom-color: #4F46E5;
    }

    [data-testid="stAlert"] {
        border-radius: 14px;
    }

    [data-testid="stProgressBar"] > div > div {
        border-radius: 8px;
        background: linear-gradient(90deg, #4F46E5, #0D9488);
    }

    hr {
        border-color: #E2E8F0;
    }

    @keyframes workspace-enter {
        from { opacity: 0; transform: translateY(5px); }
        to { opacity: 1; transform: translateY(0); }
    }

    @media (max-width: 768px) {
        .block-container {
            padding: 1rem 1rem 2rem;
        }
        .hero-card {
            padding: 1.2rem;
        }
    }

    @media (prefers-reduced-motion: reduce) {
        *, *::before, *::after {
            animation-duration: 0.01ms !important;
            animation-iteration-count: 1 !important;
            transition-duration: 0.01ms !important;
            scroll-behavior: auto !important;
        }
    }
    </style>
    """,
    unsafe_allow_html=True,
)


def open_new_screening():
    st.session_state["workspace_view"] = "New screening"


def open_overview():
    st.session_state["workspace_view"] = "Overview"


def open_saved_results(batch_id):
    archived = st.session_state.get("screening_result_archive", {}).get(batch_id)
    if archived is None:
        st.session_state["history_notice"] = (
            "This batch's candidate details are no longer available in the current session."
        )
        return

    st.session_state["screening_results"] = pd.read_json(
        io.StringIO(archived),
        orient="split",
    )
    batch = next(
        (item for item in st.session_state.get("screening_history", [])
         if item.get("batch_id") == batch_id),
        {},
    )
    st.session_state["active_job_title"] = batch.get(
        "job_title",
        "Previous screening",
    )
    st.session_state["workspace_view"] = "New screening"


def add_chart_hover(chart):
    chart.update_layout(
        hovermode="closest",
        hoverlabel={
            "bgcolor": "#171A2B",
            "bordercolor": "#6257E8",
            "font": {"color": "#FFFFFF", "size": 13},
        },
    )
    for trace in chart.data:
        if trace.type == "histogram":
            trace.hovertemplate = (
                "Score: %{x}<br>Candidates in score band: %{y}<extra></extra>"
            )
        elif trace.type == "bar" and trace.orientation == "h":
            trace.hovertemplate = (
                "<b>%{y}</b><br>%{fullData.name}: %{x:.1f}"
                "<extra></extra>"
            )
        elif trace.type == "bar":
            trace.hovertemplate = (
                "<b>%{x}</b><br>%{fullData.name}: %{y:.1f}"
                "<extra></extra>"
            )
    return chart


def normalize_match_keys(values):
    """Normalize human-entered identifiers for case- and whitespace-safe joins."""
    return values.fillna("").astype(str).str.strip().str.casefold()


def resolve_verified_labels(results_df, labels, candidate_id_column, resume_column):
    """Resolve uploaded labels to this batch by ID first, then unique resume name."""
    resolved_labels = labels.copy()
    result_ids = results_df["Candidate_ID"].astype(str)
    id_lookup = dict(zip(normalize_match_keys(result_ids), result_ids))

    resume_lookup = {}
    duplicate_result_names = set()
    if resume_column and "Resume" in results_df.columns:
        result_names = results_df["Resume"].map(
            lambda value: Path(str(value)).name
        )
        normalized_names = normalize_match_keys(result_names)
        duplicated_names = set(
            normalized_names[normalized_names.duplicated(keep=False)].tolist()
        )
        duplicate_result_names = duplicated_names
        resume_lookup = {
            key: candidate_id
            for key, candidate_id in zip(normalized_names, result_ids)
            if key and key not in duplicated_names
        }

    resolved_ids = []
    for row_index in range(len(resolved_labels)):
        row = resolved_labels.iloc[row_index]
        candidate_id = ""
        if candidate_id_column:
            candidate_key = normalize_match_keys(
                pd.Series([row[candidate_id_column]])
            ).iloc[0]
            candidate_id = id_lookup.get(candidate_key, "")

        if not candidate_id and resume_column and resume_lookup:
            resume_name = Path(str(row[resume_column])).name
            resume_key = normalize_match_keys(pd.Series([resume_name])).iloc[0]
            candidate_id = resume_lookup.get(resume_key, "")

        resolved_ids.append(candidate_id)

    resolved_labels["Candidate_ID"] = resolved_ids
    matched_labels = resolved_labels[resolved_labels["Candidate_ID"] != ""].copy()
    if matched_labels["Candidate_ID"].duplicated().any():
        return (
            matched_labels,
            matched_labels["Candidate_ID"].nunique(),
            "Multiple label rows resolve to the same candidate. Keep one row per candidate.",
        )

    matched_count = matched_labels["Candidate_ID"].nunique()
    if matched_count:
        return matched_labels, matched_count, None

    message = (
        "No uploaded Candidate_ID or resume filename matched this screening. "
        "Use the complete results CSV from this batch as the label template; "
        "synthetic DEMO IDs only work in the synthetic demonstration."
    )
    if duplicate_result_names and not candidate_id_column:
        message = (
            "Resume filenames are not unique in this batch. Include Candidate_ID "
            "from the downloaded results CSV to match labels reliably."
        )
    return matched_labels, 0, message


def calculate_fairness_rates(
    results_df,
    demographic_df,
    group_column,
    match_column="Candidate_ID",
):
    """Calculate selection-rate gaps from explicitly supplied demographic labels."""
    labels = demographic_df.copy()
    if match_column not in labels.columns or match_column not in results_df.columns:
        return pd.DataFrame(), pd.DataFrame(), None, None

    labels["_audit_match_key"] = normalize_match_keys(labels[match_column])
    result_columns = list(dict.fromkeys(["Candidate_ID", "Category", match_column]))
    joined_source = results_df[result_columns].copy()
    joined_source["_audit_match_key"] = normalize_match_keys(
        joined_source[match_column]
    )
    labels_for_join = labels.drop(
        columns=[column for column in ("Candidate_ID", "Resume") if column != "_audit_match_key"],
        errors="ignore",
    )
    joined = joined_source.merge(
        labels_for_join,
        on="_audit_match_key",
        how="inner",
        suffixes=("", "_label"),
    ).drop(columns="_audit_match_key")

    invalid_labels = {
        "", "unknown", "unspecified", "not disclosed",
        "prefer not to say", "not provided", "not independently assessed",
        "not assessed", "n/a", "na", "nan", "none",
    }
    joined[group_column] = joined[group_column].astype(str).str.strip()
    joined = joined[~joined[group_column].str.lower().isin(invalid_labels)]
    if joined.empty:
        return joined, pd.DataFrame(), None, None

    rates = joined.groupby(group_column).agg(
        Total=("Candidate_ID", "count"),
        Shortlisted=(
            "Category",
            lambda values: int((values == "Shortlisted").sum()),
        ),
    )
    rates["Shortlist_Rate"] = rates["Shortlisted"] / rates["Total"] * 100
    parity_gap = rates["Shortlist_Rate"].max() - rates["Shortlist_Rate"].min()

    opportunity_gap = None
    if "Qualified" in joined.columns:
        qualified_values = joined["Qualified"].astype(str).str.strip().str.lower()
        joined["Qualified_Valid"] = qualified_values.isin(
            {"true", "yes", "1", "qualified"}
        )
        qualified = joined[joined["Qualified_Valid"]]
        qualified_counts = qualified.groupby(group_column).agg(
            Qualified_Total=("Candidate_ID", "count"),
            Qualified_Shortlisted=(
                "Category",
                lambda values: int((values == "Shortlisted").sum()),
            ),
        )
        qualified_counts["Equal_Opportunity_Rate"] = (
            qualified_counts["Qualified_Shortlisted"]
            / qualified_counts["Qualified_Total"]
            * 100
        )
        if len(qualified_counts) >= 2:
            rates = rates.join(qualified_counts, how="left")
            opportunity_rates = qualified_counts["Equal_Opportunity_Rate"]
            opportunity_gap = (
                opportunity_rates.max() - opportunity_rates.min()
            )

    return joined, rates, float(parity_gap), opportunity_gap


def build_fairness_export(results_df, verified_labels=None):
    """Export screening facts and only explicitly supplied audit labels."""
    export_df = results_df.copy().astype(object)
    export_df = export_df.where(pd.notna(export_df), "Not provided")
    for column in export_df.columns:
        export_df[column] = export_df[column].apply(
            lambda value: (
                "Not provided"
                if isinstance(value, str) and not value.strip()
                else value
            )
        )
    export_df["Gender"] = "Not provided"
    export_df["Age_Group"] = "Not provided"
    export_df["Qualified"] = "Not independently assessed"
    export_df["Qualification_Source"] = "Not provided"

    if verified_labels is not None and not verified_labels.empty:
        normalized_columns = {
            re.sub(r"[^a-z]", "", column.lower()): column
            for column in verified_labels.columns
        }
        id_column = normalized_columns.get("candidateid")
        if id_column:
            labels = verified_labels.copy()
            labels["_Candidate_ID"] = labels[id_column].astype(str).str.strip()
            label_lookup = labels.drop_duplicates("_Candidate_ID").set_index(
                "_Candidate_ID"
            )
            candidate_ids = export_df["Candidate_ID"].astype(str).str.strip()
            source_columns = {
                "Gender": normalized_columns.get("gender"),
                "Age_Group": normalized_columns.get("agegroup"),
                "Qualified": normalized_columns.get("qualified"),
                "Qualification_Source": normalized_columns.get("qualificationsource"),
            }
            for output_column, source_column in source_columns.items():
                if not source_column:
                    continue
                mapped = candidate_ids.map(label_lookup[source_column])
                valid = mapped.notna() & mapped.astype(str).str.strip().ne("")
                export_df.loc[valid, output_column] = mapped.loc[valid].astype(str)

            qualified_source = source_columns["Qualified"]
            qualification_source = source_columns["Qualification_Source"]
            if qualified_source and not qualification_source:
                has_qualified_label = export_df["Qualified"].isin(
                    {"yes", "no", "true", "false", "1", "0", "qualified", "not qualified"}
                )
                export_df.loc[has_qualified_label, "Qualification_Source"] = (
                    "Recruiter-provided CSV"
                )

    export_df = export_df.where(pd.notna(export_df), "Not provided")
    for column in export_df.columns:
        export_df[column] = export_df[column].apply(
            lambda value: (
                "Not provided"
                if isinstance(value, str) and not value.strip()
                else value
            )
        )
    return export_df


with st.sidebar:
    st.markdown(":material/blur_on: **TalentLens**")
    st.caption("Recruitment intelligence")
    workspace_view = st.radio(
        "Workspace",
        ["Overview", "New screening", "Screening history", "Fairness audit"],
        key="workspace_view",
        label_visibility="collapsed",
    )
    st.caption("Your screening data stays in this session.")


screening_history = st.session_state.get("screening_history", [])
latest_results = st.session_state.get("screening_results")
role_title = st.session_state.get("active_job_title", "No role selected")

if workspace_view == "Screening history":
    st.title("Screening history", icon=":material/history:")
    st.caption("Completed batches for this browser session")
    if screening_history:
        for session in screening_history[:20]:
            history_columns = st.columns([2.4, 1, 1, 1, 1.2])
            history_columns[0].markdown(
                f"**{session['job_title']}**  \n{session['created_at']}"
            )
            history_columns[1].caption(f"{session['total']} resumes")
            history_columns[2].caption(f"{session['shortlisted']} shortlisted")
            history_columns[3].caption(f"{session['under_review']} to review")
            if session.get("batch_id"):
                history_columns[4].button(
                    "View results",
                    key=f"history_view_{session['batch_id']}",
                    icon=":material/arrow_forward:",
                    on_click=open_saved_results,
                    args=(session["batch_id"],),
                )
            else:
                history_columns[4].caption("Results unavailable")
    else:
        st.info("Completed screening batches will appear here.")
    st.button(
        "Start a screening",
        type="primary",
        icon=":material/add:",
        on_click=open_new_screening,
    )
    st.stop()


if workspace_view == "Fairness audit":
    st.title("Fairness audit", icon=":material/balance:")
    st.caption(
        "Review selection-rate differences using verified labels. "
        "A disparity is a signal for investigation, not proof of bias or fairness."
    )

    if latest_results is None or latest_results.empty:
        st.info("Complete a screening before auditing group outcomes.")
        st.stop()

    # # st.caption(
    # #     "The CSV includes every screening result field. Gender/Age_Group are not "
    # #     "inferred from resumes; replace “Not provided” only with verified labels. "
    # #     "Qualified is kept separate from the model's screening category and must "
    # #     "be independently assessed before using equal-opportunity analysis."
    # )
    labels_upload = st.file_uploader(
        "Upload verified labels",
        type=["csv"],
        key="fairness_labels_upload",
        help="Required: Candidate_ID and Gender and/or Age_Group. Qualified is optional.",
    )

    show_fairness_demo = st.toggle(
        "Show synthetic fairness demonstration",
        help="Displays clearly synthetic records separate from real candidates.",
    )
    if show_fairness_demo:
        demo_rows = [
            ("DEMO-W01", "Woman", "25-34", "yes", "Shortlisted", 86),
            ("DEMO-W02", "Woman", "25-34", "yes", "Shortlisted", 82),
            ("DEMO-W03", "Woman", "35-44", "yes", "Shortlisted", 78),
            ("DEMO-W04", "Woman", "45-54", "yes", "Under Review", 58),
            ("DEMO-W05", "Woman", "35-44", "no", "Rejected", 22),
            ("DEMO-M01", "Man", "25-34", "yes", "Shortlisted", 75),
            ("DEMO-M02", "Man", "35-44", "yes", "Under Review", 61),
            ("DEMO-M03", "Man", "35-44", "yes", "Rejected", 43),
            ("DEMO-M04", "Man", "45-54", "yes", "Rejected", 34),
            ("DEMO-M05", "Man", "55-64", "no", "Rejected", 18),
        ]
        demo_data = pd.DataFrame(
            demo_rows,
            columns=[
                "Candidate_ID",
                "Gender",
                "Age_Group",
                "Qualified",
                "Category",
                "Final_Score",
            ],
        )
        st.warning(
            "Synthetic demonstration only. These DEMO records are invented and "
            "are not linked to uploaded resumes or real candidates. Do not use "
            "these values for a hiring decision.",
            icon=":material/science:",
        )
        st.download_button(
            "Download synthetic demo CSV",
            data=demo_data.to_csv(index=False).encode("utf-8"),
            file_name="synthetic_fairness_demo_only.csv",
            mime="text/csv",
            icon=":material/download:",
        )
        demo_results = demo_data[["Candidate_ID", "Category"]]
        demo_labels = demo_data[
            ["Candidate_ID", "Gender", "Age_Group", "Qualified"]
        ]
        _, demo_gender_rates, demo_parity_gap, demo_opportunity_gap = (
            calculate_fairness_rates(
                demo_results,
                demo_labels,
                "Gender",
            )
        )
        demo_metrics = st.columns(2)
        demo_metrics[0].metric(
            "Illustrative selection-rate gap",
            f"{demo_parity_gap:.1f} pp",
        )
        demo_metrics[1].metric(
            "Illustrative qualified-selection gap",
            f"{demo_opportunity_gap:.1f} pp",
        )
        demo_chart = go.Figure(
            go.Bar(
                name="Selection rate",
                x=demo_gender_rates.index.tolist(),
                y=demo_gender_rates["Shortlist_Rate"].tolist(),
                marker_color=["#6257E8", "#4D79C7"],
                text=[
                    f"{rate:.0f}% · n={total}"
                    for rate, total in zip(
                        demo_gender_rates["Shortlist_Rate"],
                        demo_gender_rates["Total"],
                    )
                ],
                textposition="outside",
            )
        )
        demo_chart.update_layout(
            height=320,
            yaxis_title="Shortlisted (%)",
            yaxis_range=[0, 100],
            showlegend=False,
        )

        demo_qualified = demo_data[demo_data["Qualified"] == "yes"]
        demo_qualified_rates = demo_qualified.groupby("Gender").agg(
            Qualified_Total=("Candidate_ID", "count"),
            Qualified_Shortlisted=(
                "Category",
                lambda values: int((values == "Shortlisted").sum()),
            ),
        )
        demo_qualified_rates["Rate"] = (
            demo_qualified_rates["Qualified_Shortlisted"]
            / demo_qualified_rates["Qualified_Total"]
            * 100
        )
        demo_opportunity_chart = go.Figure(
            go.Bar(
                name="Qualified selection rate",
                x=demo_qualified_rates.index.tolist(),
                y=demo_qualified_rates["Rate"].tolist(),
                marker_color=["#8CBDB5", "#B98B43"],
                text=[
                    f"{rate:.0f}% · n={total}"
                    for rate, total in zip(
                        demo_qualified_rates["Rate"],
                        demo_qualified_rates["Qualified_Total"],
                    )
                ],
                textposition="outside",
            )
        )
        demo_opportunity_chart.update_layout(
            height=320,
            yaxis_title="Qualified candidates shortlisted (%)",
            yaxis_range=[0, 100],
            showlegend=False,
        )
        selection_graph, opportunity_graph = st.columns(2)
        with selection_graph:
            st.caption("Selection rate across all synthetic candidates")
            st.plotly_chart(
                add_chart_hover(demo_chart),
                alt="Synthetic shortlist rates across illustrative gender groups",
            )
        with opportunity_graph:
            st.caption("Selection rate among synthetic qualified candidates")
            st.plotly_chart(
                add_chart_hover(demo_opportunity_chart),
                alt="Synthetic qualified-candidate shortlist rates by illustrative group",
            )

    warning_gap = 10
    critical_gap = 20
    minimum_group_size = 5
    # st.caption(
    #     "Review indicators use fixed bands: yellow at 10 percentage points, red at 20; "
    #     "groups with fewer than five candidates are marked insufficient."
    # )
    resume_categories = latest_results.get(
        "Resume_Category",
        pd.Series(["Unspecified"] * len(latest_results)),
    ).fillna("Unspecified").astype(str).unique().tolist()
    selected_resume_category = st.selectbox(
        "Resume/job category",
        ["All categories", *sorted(resume_categories)],
        help="Folder-derived categories are used when present in the screening results.",
    )
    audit_results = latest_results
    if selected_resume_category != "All categories":
        audit_results = latest_results[
            latest_results.get(
                "Resume_Category",
                pd.Series(["Unspecified"] * len(latest_results), index=latest_results.index),
            ).fillna("Unspecified").astype(str)
            == selected_resume_category
        ]

    verified_labels = None
    if labels_upload is None:
        st.info(
            "Reliable demographic group labels are not available in the current resumes. "
            "Gender- and age-based fairness metrics cannot be calculated from these records alone.",
            title="Demographic fairness audit unavailable",
        )
    else:
        try:
            labels = pd.read_csv(labels_upload)
            normalized = {
                re.sub(r"[^a-z]", "", column.lower()): column
                for column in labels.columns
            }
            candidate_id_column = normalized.get("candidateid")
            resume_column = normalized.get("resume")
            group_columns = {
                "Gender": normalized.get("gender"),
                "Age group": normalized.get("agegroup"),
            }
            if not candidate_id_column and not resume_column:
                st.error("The CSV needs a Candidate_ID or Resume column to match this screening.")
            elif not any(group_columns.values()):
                st.error("Add a Gender and/or Age_Group column to calculate group rates.")
            else:
                labels, matched_count, match_error = resolve_verified_labels(
                    audit_results,
                    labels,
                    candidate_id_column,
                    resume_column,
                )
                if match_error:
                    st.error(match_error)
                else:
                    verified_labels = labels
                    st.caption(
                        f"Verified labels matched to {matched_count} of "
                        f"{len(audit_results)} screened candidates in this category."
                    )
                    if matched_count < len(audit_results):
                        st.warning(
                            "Some screened candidates have no matching demographic label; "
                            "they are excluded from group rates."
                        )

                    for dimension, group_column in group_columns.items():
                        if not group_column:
                            continue
                        joined, rates, _, opportunity_gap = calculate_fairness_rates(
                            audit_results,
                            labels,
                            group_column,
                        )
                        if rates.empty:
                            st.info(f"No valid {dimension.lower()} labels matched this screening.")
                            continue

                        sufficiently_sized = rates[
                            rates["Total"] >= minimum_group_size
                        ]
                        st.subheader(f"{dimension} selection rates")
                        st.caption(
                            "Demographic parity compares the share shortlisted in each group. "
                            "It does not account for qualification or explain why rates differ."
                        )
                        if len(sufficiently_sized) < 2:
                            st.warning(
                                f"Insufficient sample size: fewer than two {dimension.lower()} "
                                f"groups have at least {minimum_group_size} labeled candidates."
                            )
                        else:
                            parity_gap = (
                                sufficiently_sized["Shortlist_Rate"].max()
                                - sufficiently_sized["Shortlist_Rate"].min()
                            )
                            if parity_gap >= critical_gap:
                                st.error(
                                    f"Potential disparity: selection-rate gap is "
                                    f"{parity_gap:.1f} percentage points. Review group context and criteria."
                                )
                            elif parity_gap >= warning_gap:
                                st.warning(
                                    f"Selection-rate gap is {parity_gap:.1f} percentage points; "
                                    "consider further review."
                                )
                            else:
                                st.success(
                                    f"Selection-rate gap is {parity_gap:.1f} percentage points "
                                    "for groups meeting the sample threshold."
                                )

                            selection_chart = go.Figure(
                                go.Bar(
                                    x=sufficiently_sized.index.tolist(),
                                    y=sufficiently_sized["Shortlist_Rate"].tolist(),
                                    marker_color="#6257E8",
                                    text=[
                                        f"{rate:.1f}% · n={total}"
                                        for rate, total in zip(
                                            sufficiently_sized["Shortlist_Rate"],
                                            sufficiently_sized["Total"],
                                        )
                                    ],
                                    textposition="outside",
                                )
                            )
                            selection_chart.update_layout(
                                height=320,
                                yaxis_title="Shortlisted (%)",
                                yaxis_range=[0, 100],
                                showlegend=False,
                            )
                            st.plotly_chart(
                                add_chart_hover(selection_chart),
                                alt=f"Shortlist rate by explicitly provided {dimension.lower()} groups",
                            )

                        if opportunity_gap is None:
                            st.info(
                                "Equal opportunity is unavailable without verified Qualified yes/no "
                                "labels and at least two groups with qualified candidates."
                            )
                        else:
                            opportunity_rates = rates.dropna(
                                subset=["Equal_Opportunity_Rate"]
                            )
                            opportunity_rates = opportunity_rates[
                                opportunity_rates["Qualified_Total"] >= minimum_group_size
                            ]
                            if len(opportunity_rates) < 2:
                                st.warning(
                                    "Insufficient qualified-candidate sample size to compare "
                                    "equal opportunity across two groups."
                                )
                            else:
                                opportunity_gap = (
                                    opportunity_rates["Equal_Opportunity_Rate"].max()
                                    - opportunity_rates["Equal_Opportunity_Rate"].min()
                                )
                                st.markdown("**Equal opportunity**")
                                st.caption(
                                    "Among candidates labeled Qualified, this compares the share "
                                    "shortlisted in each group."
                                )
                                if opportunity_gap >= critical_gap:
                                    st.error(
                                        f"Potential qualified-selection disparity: "
                                        f"{opportunity_gap:.1f} percentage points."
                                    )
                                elif opportunity_gap >= warning_gap:
                                    st.warning(
                                        f"Qualified-selection gap is {opportunity_gap:.1f} "
                                        "percentage points; consider further review."
                                    )
                                else:
                                    st.success(
                                        f"Qualified-selection gap is {opportunity_gap:.1f} "
                                        "percentage points for groups meeting the sample threshold."
                                    )
                                opportunity_chart = go.Figure(
                                    go.Bar(
                                        x=opportunity_rates.index.tolist(),
                                        y=opportunity_rates[
                                            "Equal_Opportunity_Rate"
                                        ].tolist(),
                                        marker_color="#4D79C7",
                                        text=[
                                            f"{rate:.1f}% · n={total}"
                                            for rate, total in zip(
                                                opportunity_rates[
                                                    "Equal_Opportunity_Rate"
                                                ],
                                                opportunity_rates["Qualified_Total"],
                                            )
                                        ],
                                        textposition="outside",
                                    )
                                )
                                opportunity_chart.update_layout(
                                    height=320,
                                    yaxis_title="Qualified candidates shortlisted (%)",
                                    yaxis_range=[0, 100],
                                    showlegend=False,
                                )
                                st.plotly_chart(
                                    add_chart_hover(opportunity_chart),
                                    alt=(
                                        "Shortlist rate among explicitly labeled "
                                        "qualified candidates by group"
                                    ),
                                )

        except Exception as exc:
            verified_labels = None
            st.error(f"Could not read the verified-label CSV: {exc}")

    audit_export = build_fairness_export(
        latest_results,
        verified_labels,
    )
    st.download_button(
        "Download complete screening and audit CSV",
        data=audit_export.to_csv(index=False).encode("utf-8"),
        file_name="screening_fairness_audit.csv",
        mime="text/csv",
        icon=":material/download:",
    )
    st.stop()


if workspace_view == "Overview":
    current_hour = datetime.now().hour
    greeting = "Good morning" if current_hour < 12 else (
        "Good afternoon" if current_hour < 18 else "Good evening"
    )

    with st.container(key="overview-hero"):
        st.caption("TALENTLENS  /  RECRUITMENT INTELLIGENCE")
        st.title("Every resume. One clearer path to great talent.")
        st.markdown(
            "Turn resume batches into structured, explainable candidate "
            "insights. Find job-related evidence sooner and keep the final "
            "decision human-led."
        )
        with st.container(horizontal=True, vertical_alignment="center"):
            st.button(
                "Start screening",
                type="primary",
                icon=":material/arrow_forward:",
                on_click=open_new_screening,
            )
            st.markdown("[Explore the platform ↓](#workflow)")
        st.caption("AI-assisted screening. Human-led decisions.")

    if screening_history:
        overview_counts = {
            "total": sum(item["total"] for item in screening_history),
            "shortlisted": sum(item["shortlisted"] for item in screening_history),
            "under_review": sum(item["under_review"] for item in screening_history),
            "rejected": sum(item["rejected"] for item in screening_history),
        }
    elif latest_results is not None:
        overview_counts = {
            "total": len(latest_results),
            "shortlisted": int((latest_results["Category"] == "Shortlisted").sum()),
            "under_review": int((latest_results["Category"] == "Under Review").sum()),
            "rejected": int((latest_results["Category"] == "Rejected").sum()),
        }
    else:
        overview_counts = {
            "total": 0,
            "shortlisted": 0,
            "under_review": 0,
            "rejected": 0,
        }

    metric_columns = st.columns(4)
    metric_columns[0].metric(
        "Resumes processed",
        overview_counts["total"],
        border=True,
        height="stretch",
    )
    metric_columns[1].metric(
        "Shortlisted",
        overview_counts["shortlisted"],
        border=True,
        height="stretch",
    )
    metric_columns[2].metric(
        "Awaiting review",
        overview_counts["under_review"],
        border=True,
        height="stretch",
    )
    metric_columns[3].metric(
        "Rejected",
        overview_counts["rejected"],
        border=True,
        height="stretch",
    )

    role_col, next_step_col = st.columns([1.2, 1])
    with role_col:
        st.subheader("Current role")
        if latest_results is None:
            st.info(
                "Your workspace is ready. Start a screening to add a job and review its candidates.",
                icon=":material/work_outline:",
            )
        else:
            st.markdown(f"**{role_title}**")
            st.caption(
                f"Latest batch: {len(latest_results)} resumes parsed and ranked."
            )
            st.progress(1.0, text="Latest screening complete")

    with next_step_col:
        st.subheader("Next step")
        if overview_counts["under_review"]:
            st.markdown(
                f"**{overview_counts['under_review']} candidates** are awaiting recruiter review."
            )
            st.caption("Open the screening workspace to inspect candidate evidence and reasons.")
        elif overview_counts["total"]:
            st.markdown("**You're all caught up on this role.**")
            st.caption("Start another screening whenever a new role is ready.")
        else:
            st.markdown("**Prepare a job description and resume folder.**")
            st.caption("The screening workspace supports pasted JDs, URLs, PDF, and DOCX.")

    st.subheader("Recent screening sessions")
    st.caption("Completed batches in this browser session")
    if screening_history:
        for session in screening_history[:5]:
            session_columns = st.columns([2.4, 1, 1, 1])
            session_columns[0].markdown(
                f"**{session['job_title']}**  \n{session['created_at']}"
            )
            session_columns[1].caption(f"{session['total']} resumes")
            session_columns[2].caption(f"{session['shortlisted']} shortlisted")
            session_columns[3].caption(f"{session['under_review']} to review")
    else:
        st.caption("No completed screenings in this session yet.")

    with st.container(border=True):
        st.markdown("**Efficiency matters. Fairness does too.**")
        st.caption(
            "The Fairness audit can compare demographic parity using recruiter-verified "
            "group labels, and equal opportunity when verified qualification outcomes "
            "are provided. Without those labels, metrics remain unavailable. Any "
            "disparity is a signal for human investigation, not proof of bias or fairness."
        )

    st.markdown('<div id="workflow"></div>', unsafe_allow_html=True)
    st.header("From resume overload to hiring clarity")
    workflow_columns = st.columns(3)
    workflow_steps = [
        (
            "01",
            "Define the role",
            "Paste a job description or review text retrieved from a LinkedIn posting.",
        ),
        (
            "02",
            "Discover relevant talent",
            "Upload a resume folder and compare skills, experience, and education against the role.",
        ),
        (
            "03",
            "Review with context",
            "Inspect the ranking, component scores, matched skills, and decision rationale before deciding.",
        ),
    ]
    for column, (step_number, title, description) in zip(
        workflow_columns,
        workflow_steps,
    ):
        with column:
            st.caption(step_number)
            st.subheader(title)
            st.write(description)

    st.header("Built for informed, human-led decisions")
    trust_columns = st.columns(3)
    trust_columns[0].markdown("**Explainable recommendations**")
    trust_columns[0].caption(
        "Candidate scores include skill, experience, and education components."
    )
    trust_columns[1].markdown("**Job-related evidence**")
    trust_columns[1].caption(
        "The screening compares extracted resume information with the selected role."
    )
    trust_columns[2].markdown("**Human oversight**")
    trust_columns[2].caption(
        "Results are review aids; recruiters remain responsible for hiring decisions."
    )

    st.stop()


# --------------------------------------------------
# 1. EXTRACT TEXT FROM PDF AND DOCX
# --------------------------------------------------

def extract_resume_text(uploaded_file):
    """Extract text from an uploaded PDF or DOCX file."""
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if filename.endswith(".pdf"):
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages = [
                page.extract_text() or ""
                for page in pdf.pages
            ]
        return "\n".join(pages).strip()

    if filename.endswith(".docx"):
        document = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs]
        return "\n".join(paragraphs).strip()

    return ""


# --------------------------------------------------
# 2. CLASSIFY CANDIDATES
# --------------------------------------------------

def classify_candidate(
    final_score,
    rank,
    required_experience_met,
    classification_mode,
    shortlist_value,
):
    review_score_floor = 50
    if not required_experience_met:
        return "Rejected"

    if classification_mode == "Score thresholds":
        if final_score >= shortlist_value:
            return "Shortlisted"
    elif rank <= shortlist_value:
        return "Shortlisted"

    if final_score >= review_score_floor:
        return "Under Review"
    return "Rejected"


def get_gate_score(report, gate_name):
    """Get a component score from the gate report."""
    gate = report.get("gate_results", {}).get(gate_name, {})
    return float(gate.get("score_pct", 0)) * 100


def get_skill_match_score(jd, report):
    """Return the percentage of listed JD skills matched by the resume."""
    requirements = jd.get("requirements", {}).get("skills", {})
    required_skills = [
        item.get("skill", "")
        for group in ("must_have", "preferred")
        for item in requirements.get(group, [])
        if item.get("skill")
    ]

    if not required_skills:
        return 100.0

    matched_skills = {
        re.sub(r"\s*\(partial\)$", "", skill, flags=re.IGNORECASE).lower()
        for skill in report.get("matched_skills", [])
    }
    matched_count = sum(
        skill.lower() in matched_skills for skill in required_skills
    )
    return matched_count / len(required_skills) * 100


def get_education_match_score(jd, profile):
    """Give partial credit for a related but lower-level degree."""
    required = jd.get("requirements", {}).get("education", {}).get(
        "required_degrees", []
    )
    candidate = profile.get("education", {}).get("degrees", [])
    if not required:
        return 100.0
    if not candidate:
        return 0.0

    degree_levels = {
        "diploma": 1,
        "associate": 1,
        "bachelor": 2,
        "master": 3,
        "doctor": 4,
        "phd": 4,
    }

    def degree_level(degree):
        degree = degree.lower()
        for label, level in degree_levels.items():
            if label in degree:
                return level
        return 0

    best_score = 0.0
    for required_degree in required:
        required_level = degree_level(required_degree)
        for candidate_degree in candidate:
            candidate_level = degree_level(candidate_degree)
            if required_degree.lower() in candidate_degree.lower():
                score = 100.0
            elif candidate_level > required_level > 0:
                score = 85.0
            elif candidate_level and required_level:
                score = 55.0
            else:
                score = 0.0
            best_score = max(best_score, score)
    return best_score


def get_resume_category(uploaded_file):
    """Use an uploaded subfolder as the category when directory paths exist."""
    relative_path = getattr(uploaded_file, "relative_path", None)
    path = Path(str(relative_path or uploaded_file.name).replace("\\", "/"))
    parent = path.parent.name
    return parent if parent not in ("", ".") else "Unspecified"


def count_skills_by_candidate(values):
    """Count how many resumes mention each skill, counting each once per resume."""
    counts = Counter()
    for value in values.fillna(""):
        skills = {
            re.sub(r"\s*\(partial\)$", "", skill, flags=re.IGNORECASE).strip()
            for skill in str(value).split(",")
        }
        counts.update(skill for skill in skills if skill and skill.lower() != "none")
    return counts.most_common(10)


def get_classification_reason(candidate, classification_mode, shortlist_value):
    """Explain the decision using the candidate's measured component scores."""
    skill_score = candidate["Skill_Match_%"]
    experience_score = candidate["Experience_Match_%"]
    education_score = candidate["Education_Match_%"]
    final_score = candidate["Final_Score"]
    matched = candidate["Matched_Skills"] or "none"
    missing = candidate["Missing_Skills"] or "none"
    evidence = (
        f"Skills {skill_score:.1f}%, experience {experience_score:.1f}% "
        f"({candidate['Experience_Years']:.1f} years parsed), and education "
        f"{education_score:.1f}%. Matched skills: {matched}. "
        f"Missing skills: {missing}."
    )

    if not candidate["Required_Experience_Met"]:
        return f"Rejected: {candidate['Experience_Reason']} {evidence}"

    if candidate["Category"] == "Shortlisted":
        if classification_mode == "Score thresholds":
            basis = f"final score {final_score:.1f} met the {shortlist_value} shortlist threshold"
        else:
            basis = f"rank {int(candidate['Rank'])} was within the top {shortlist_value}"
        return f"Shortlisted: {basis}. {evidence}"

    if candidate["Category"] == "Under Review":
        if classification_mode == "Score thresholds":
            basis = (
                f"final score {final_score:.1f} was below the {shortlist_value} "
                "shortlist threshold but met the 50-point review floor"
            )
        else:
            basis = (
                f"rank {int(candidate['Rank'])} was outside the top {shortlist_value} "
                "shortlist ranks, while the score met the 50-point review floor"
            )
        return f"Under Review: {basis}. {evidence}"

    low_components = []
    if skill_score < 50:
        low_components.append(f"skills ({skill_score:.1f}%)")
    if experience_score < 50:
        low_components.append(f"experience ({experience_score:.1f}%)")
    if education_score < 50:
        low_components.append(f"education ({education_score:.1f}%)")
    weak_reason = ", ".join(low_components) or "overall score"
    return (
        f"Rejected: final score {final_score:.1f} was below the 50-point "
        f"review floor; low match in {weak_reason}. {evidence}"
    )


# --------------------------------------------------
# 3. JOB DESCRIPTION INPUT
# --------------------------------------------------

st.title("New screening", icon=":material/person_search:")
st.caption("Set the role requirements, add a resume batch, and review the ranked candidates.")

st.subheader("1. Choose a job", icon=":material/work_outline:")

jd_method = st.radio(
    "Choose how to provide the job description",
    ["Paste job description", "Job URL"],
    horizontal=True,
)

job_title = st.text_input(
    "Job title",
    placeholder="e.g. Data Analyst",
)

jd_text = ""

if jd_method == "Paste job description":
    jd_text = st.text_area(
        "Paste the full job description",
        height=220,
        placeholder=(
            "Include required skills, education, minimum experience "
            "and preferred qualifications."
        ),
    )

else:
    job_url = st.text_input(
        "LinkedIn job posting URL",
        placeholder="https://www.linkedin.com/jobs/view/...",
    )

    if st.button("Fetch job description"):
        if not job_url.startswith("https://"):
            st.error("Please enter a valid HTTPS URL.")
        else:
            with st.spinner("Attempting to retrieve the job description..."):
                scraped = scrape_job_url(job_url)

            if scraped and len(scraped.strip()) > 100:
                st.session_state["scraped_jd_text"] = scraped
                st.success("Page text retrieved. Review it before screening.")
            else:
                st.warning(
                    "The page could not be retrieved reliably. "
                    "LinkedIn may restrict automated access. "
                    "Paste the job description manually instead."
                )

    jd_text = st.session_state.get("scraped_jd_text", "")

    if jd_text:
        jd_text = st.text_area(
            "Review or edit the retrieved job description",
            value=jd_text,
            height=220,
            key="reviewed_jd_text",
        )


# --------------------------------------------------
# 4. RESUME UPLOAD
# --------------------------------------------------

st.subheader("2. Upload resumes", icon=":material/drive_folder_upload:")

uploaded_resumes = st.file_uploader(
    "Select a resume folder or multiple PDF/DOCX files",
    type=["pdf", "docx"],
    accept_multiple_files="directory",
    help="Choose a folder to include its PDF and DOCX files, or add files individually.",
)

# --------------------------------------------------
# 5. SCREENING
# --------------------------------------------------

st.subheader("3. Review criteria and screen", icon=":material/tune:")

classification_mode = st.radio(
    "Choose how candidates are classified",
    ["Score thresholds", "Rank ranges"],
    horizontal=True,
)

if classification_mode == "Score thresholds":
    shortlist_value = st.number_input(
        "Shortlist score threshold",
        min_value=50,
        max_value=100,
        value=70,
        help="Candidates at or above this value are shortlisted.",
    )
    st.caption(
        "Below the shortlist score, candidates scoring 50 or higher are "
        "under review; lower scores are rejected."
    )
else:
    shortlist_value = st.number_input(
        "Number of candidates to shortlist",
        min_value=1,
        max_value=10000,
        value=50,
        help="The top candidates by final score, through this rank, are shortlisted.",
    )
    st.caption(
        "Outside the shortlist, candidates scoring 50 or higher are under "
        "review; lower scores are rejected. Candidates failing required "
        "experience are rejected regardless of rank."
    )

if st.button("Screen and Rank Candidates", type="primary"):

    if not jd_text or not jd_text.strip():
        st.error("Provide a job description before screening.")

    elif not uploaded_resumes:
        st.error("Upload at least one PDF or DOCX resume.")

    else:
        with st.spinner("Parsing the JD and evaluating resumes..."):

            jd = parse_job_description(
                jd_text,
                role_title=job_title or "Unspecified Role",
                source_url=(
                    job_url
                    if jd_method == "LinkedIn job URL"
                    else None
                ),
            )

            profiles = []
            raw_texts = []
            filenames = []
            errors = []

            for index, file in enumerate(uploaded_resumes, start=1):
                try:
                    text = extract_resume_text(file)
                    resume_category = get_resume_category(file)

                    if not text:
                        errors.append(
                            f"{file.name}: no readable text found"
                        )
                        continue

                    profile = parse_resume(
                        resume_id=f"UPLOAD_{index}",
                        raw_text=text,
                        category=resume_category,
                    )

                    profiles.append(profile)
                    raw_texts.append(text)
                    filenames.append(file.name)

                except Exception as exc:
                    errors.append(f"{file.name}: {exc}")

            if not profiles:
                st.error(
                    "No readable resumes were extracted. "
                    "Check that the PDFs contain selectable text."
                )

            else:
                gate_reports = [
                    run_candidate_cascade(
                        jd,
                        profile,
                        jd.get("gate_configuration", {}).get("gate_order"),
                        jd.get("gate_configuration", {}).get("gate_weights"),
                        stop_on_failure=False,
                    )
                    for profile in profiles
                ]

                semantic_available = True
                try:
                    scorer = SemanticScorer(
                        model_name="all-MiniLM-L6-v2"
                    )

                    semantic_results = scorer.compute_similarity_scores(
                        jd_text=jd_text,
                        resume_texts=raw_texts,
                        candidate_ids=[
                            p["candidate_id"] for p in profiles
                        ],
                    )

                    semantic_map = {
                        item["candidate_id"]: item["semantic_score"]
                        for item in semantic_results
                    }
                except Exception:
                    semantic_available = False
                    semantic_map = {}
                    st.warning(
                        "Semantic similarity is unavailable; candidates will "
                        "be ranked using skills, education, and experience."
                    )

                rows = []

                for profile, report, filename in zip(
                    profiles, gate_reports, filenames
                ):
                    cid = profile["candidate_id"]

                    skill_score = get_skill_match_score(jd, report)
                    education_score = get_education_match_score(jd, profile)
                    experience_score = get_gate_score(
                        report, "experience"
                    )

                    semantic_score = semantic_map.get(cid, 0.0)

                    required_experience_met = (
                        report.get("gate_results", {})
                        .get("experience", {})
                        .get("status") == "PASS"
                    )
                    relevant_experience_score = (
                        experience_score * skill_score / 100
                    )
                    final_score = round(
                        skill_score * 0.45
                        + relevant_experience_score * 0.40
                        + education_score * 0.15,
                        2,
                    )
                    if not required_experience_met:
                        final_score = min(final_score, 15.0)

                    actual_experience = float(
                        profile.get("experience", {}).get("years", 0)
                    )

                    matched = sorted(set(report.get("matched_skills", [])))
                    missing = sorted(set(report.get("missing_skills", [])))

                    rows.append({
                        "Rank": 0,
                        "Candidate_ID": cid,
                        "Resume": filename,
                        "Resume_Category": profile.get(
                            "category", "Unspecified"
                        ),
                        "Skill_Match_%": round(skill_score, 2),
                        "Education_Match_%": round(education_score, 2),
                        "Experience_Match_%": round(experience_score, 2),
                        "Semantic_Match_%": round(semantic_score, 2),
                        "Final_Score": round(final_score, 2),
                        "Category": "",
                        "Required_Experience_Met": required_experience_met,
                        "Experience_Years": actual_experience,
                        "Matched_Skills": ", ".join(matched),
                        "Missing_Skills": ", ".join(missing),
                        "Experience_Reason": (
                            report.get("gate_results", {})
                            .get("experience", {})
                            .get("reason", "Required experience not met")
                            if not required_experience_met
                            else "Required experience was met."
                        ),
                        "Reason": "",
                    })

                results_df = pd.DataFrame(rows)

                # Sort all candidates by final score, highest first.
                results_df = results_df.sort_values(
                    ["Final_Score", "Semantic_Match_%"],
                    ascending=[False, False],
                ).reset_index(drop=True)

                results_df["Rank"] = range(1, len(results_df) + 1)
                for row_index in range(len(results_df)):
                    row = results_df.iloc[row_index]
                    results_df.at[row_index, "Category"] = classify_candidate(
                        final_score=row["Final_Score"],
                        rank=row["Rank"],
                        required_experience_met=row["Required_Experience_Met"],
                        classification_mode=classification_mode,
                        shortlist_value=shortlist_value,
                    )
                    results_df.at[row_index, "Reason"] = get_classification_reason(
                        results_df.iloc[row_index],
                        classification_mode,
                        shortlist_value,
                    )

                st.session_state["screening_results"] = results_df
                st.session_state["screening_jd"] = jd
                st.session_state["active_job_title"] = (
                    job_title.strip()
                    or jd.get("metadata", {}).get("job_title", "Unspecified role")
                )
                batch_id = uuid4().hex
                result_archive = st.session_state.setdefault(
                    "screening_result_archive", {}
                )
                result_archive[batch_id] = results_df.to_json(orient="split")
                screening_history = st.session_state.setdefault(
                    "screening_history", []
                )
                screening_history.insert(
                    0,
                    {
                        "batch_id": batch_id,
                        "job_title": st.session_state["active_job_title"],
                        "created_at": datetime.now().strftime(
                            "%b %d, %Y at %I:%M %p"
                        ),
                        "total": len(results_df),
                        "shortlisted": int(
                            (results_df["Category"] == "Shortlisted").sum()
                        ),
                        "under_review": int(
                            (results_df["Category"] == "Under Review").sum()
                        ),
                        "rejected": int(
                            (results_df["Category"] == "Rejected").sum()
                        ),
                    },
                )
                for old_session in screening_history[20:]:
                    result_archive.pop(old_session.get("batch_id"), None)
                del screening_history[20:]

                if errors:
                    st.warning(
                        "Some files could not be processed:\n\n"
                        + "\n".join(errors)
                    )

                st.success(
                    f"Screening completed for {len(results_df)} resumes."
                )


# --------------------------------------------------
# 6. DISPLAY RESULTS
# --------------------------------------------------

if "screening_results" in st.session_state:
    results_df = st.session_state["screening_results"]
    if "Resume_Category" not in results_df:
        results_df = results_df.copy()
        results_df["Resume_Category"] = "Unspecified"

    st.header("Ranked screening results", icon=":material/format_list_numbered:")

    filter_col, sort_col = st.columns([1.5, 1])
    with filter_col:
        candidate_query = st.text_input(
            "Search candidates",
            placeholder="Resume filename or candidate ID",
            key="candidate_search",
        )
    with sort_col:
        sort_choice = st.selectbox(
            "Sort candidates by",
            [
                "Final score: high to low",
                "Final score: low to high",
                "Skills match: high to low",
                "Experience match: high to low",
            ],
        )

    category_filter = st.pills(
        "Candidate status",
        ["Shortlisted", "Under Review", "Rejected"],
        selection_mode="multi",
        default=["Shortlisted", "Under Review", "Rejected"],
    )
    filtered_df = results_df[
        results_df["Category"].isin(category_filter)
    ]
    if candidate_query.strip():
        query = candidate_query.strip().casefold()
        matches_candidate = (
            filtered_df["Candidate_ID"].astype(str).str.casefold().str.contains(query)
            | filtered_df["Resume"].astype(str).str.casefold().str.contains(query)
        )
        filtered_df = filtered_df[matches_candidate]

    sort_columns = {
        "Final score: high to low": ("Final_Score", False),
        "Final score: low to high": ("Final_Score", True),
        "Skills match: high to low": ("Skill_Match_%", False),
        "Experience match: high to low": ("Experience_Match_%", False),
    }
    sort_column, sort_ascending = sort_columns[sort_choice]
    filtered_df = filtered_df.sort_values(
        sort_column,
        ascending=sort_ascending,
    )
    page_size = 25
    page_count = max(1, (len(filtered_df) + page_size - 1) // page_size)
    current_page = st.number_input(
        "Results page",
        min_value=1,
        max_value=page_count,
        step=1,
    )
    start_index = (current_page - 1) * page_size
    page_df = filtered_df.iloc[start_index:start_index + page_size]
    st.caption(
        f"Showing {len(page_df)} of {len(filtered_df)} matching resumes "
        f"(page {current_page} of {page_count})."
    )

    ranked_header = st.columns([0.45, 2.4, 0.8, 0.8, 0.8, 0.8, 1.0])
    for column, label in zip(
        ranked_header,
        ["Rank", "Resume", "Skills", "Education", "Experience", "Final", "Category"],
    ):
        column.markdown(f"**{label}**")
    for _, row in page_df.iterrows():
        columns = st.columns([0.45, 2.4, 0.8, 0.8, 0.8, 0.8, 1.0])
        columns[0].write(row["Rank"])
        columns[1].write(row["Resume"])
        columns[2].write(f"{row['Skill_Match_%']:.0f}%")
        columns[3].write(f"{row['Education_Match_%']:.0f}%")
        columns[4].write(f"{row['Experience_Match_%']:.0f}%")
        columns[5].write(f"{row['Final_Score']:.1f}")
        columns[6].write(row["Category"])

    st.subheader("Candidate details", icon=":material/person:")
    selected_resume = st.selectbox(
        "Choose a resume to inspect",
        results_df["Resume"].tolist(),
    )
    candidate = results_df[results_df["Resume"] == selected_resume].iloc[0]
    with st.container(border=True):
        detail_header, detail_status = st.columns([3, 1])
        detail_header.markdown(
            f"**{candidate['Resume']}**  ·  Rank {candidate['Rank']}"
        )
        detail_status.badge(
            candidate["Category"],
            color=(
                "green"
                if candidate["Category"] == "Shortlisted"
                else "orange"
                if candidate["Category"] == "Under Review"
                else "red"
            ),
        )
        score_cols = st.columns(4)
        score_cols[0].metric("Final score", f"{candidate['Final_Score']:.1f}/100")
        score_cols[1].metric("Skills match", f"{candidate['Skill_Match_%']:.1f}%")
        score_cols[2].metric("Experience match", f"{candidate['Experience_Match_%']:.1f}%")
        score_cols[3].metric("Education match", f"{candidate['Education_Match_%']:.1f}%")
        st.markdown(f"**Decision rationale**  \n{candidate.get('Reason', 'Not available')}")
        st.caption(f"Matched skills: {candidate['Matched_Skills'] or 'None'}")
        st.caption(f"Missing skills: {candidate['Missing_Skills'] or 'None'}")

    st.download_button(
        "Download ranked results as CSV",
        data=results_df.to_csv(index=False).encode("utf-8"),
        file_name="ranked_resume_results.csv",
        mime="text/csv",
    )

    st.header("Recruiting Insights Dashboard")

    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Total Resumes", len(results_df))
    col2.metric(
        "Shortlisted",
        int((results_df["Category"] == "Shortlisted").sum()),
    )
    col3.metric(
        "Under Review",
        int((results_df["Category"] == "Under Review").sum()),
    )
    col4.metric(
        "Rejected",
        int((results_df["Category"] == "Rejected").sum()),
    )

    status_order = ["Shortlisted", "Under Review", "Rejected"]
    status_colors = {
        "Shortlisted": "#2e8b57",
        "Under Review": "#d99a22",
        "Rejected": "#c85151",
    }
    status_counts = results_df["Category"].value_counts()
    summary_col, score_col = st.columns(2)

    with summary_col:
        st.subheader("Screening outcomes")
        outcome_chart = go.Figure(
            go.Bar(
                x=status_order,
                y=[int(status_counts.get(status, 0)) for status in status_order],
                marker_color=[status_colors[status] for status in status_order],
                text=[int(status_counts.get(status, 0)) for status in status_order],
                textposition="outside",
            )
        )
        outcome_chart.update_layout(
            height=320,
            xaxis_title="",
            yaxis_title="Resumes",
            showlegend=False,
        )
        st.plotly_chart(
            add_chart_hover(outcome_chart),
            alt="Resume counts for shortlisted, under-review, and rejected candidates",
        )

    with score_col:
        st.subheader("Final score distribution")
        score_chart = go.Figure(
            go.Histogram(
                x=results_df["Final_Score"].tolist(),
                xbins={"start": 0, "end": 100, "size": 10},
                marker_color="#337f88",
            )
        )
        score_chart.update_layout(
            height=320,
            xaxis_title="Final score",
            yaxis_title="Resumes",
            bargap=0.08,
        )
        st.plotly_chart(
            add_chart_hover(score_chart),
            alt="Distribution of final resume screening scores from zero to one hundred",
        )

    st.subheader("Resume category analysis")
    st.caption(
        "Category labels come from uploaded subfolders when available; "
        "otherwise candidates are shown as Unspecified."
    )
    category_summary = results_df.groupby("Resume_Category").agg(
        Total=("Candidate_ID", "count"),
        Shortlisted=(
            "Category",
            lambda values: int((values == "Shortlisted").sum()),
        ),
    )
    category_summary["Shortlist_Rate"] = (
        category_summary["Shortlisted"] / category_summary["Total"] * 100
    )
    category_names = category_summary.index.tolist()

    category_col, rate_col = st.columns(2)
    with category_col:
        category_chart = go.Figure()
        for status in status_order:
            counts = [
                int(
                    (
                        (results_df["Resume_Category"] == category)
                        & (results_df["Category"] == status)
                    ).sum()
                )
                for category in category_names
            ]
            category_chart.add_trace(
                go.Bar(
                    name=status,
                    x=category_names,
                    y=counts,
                    marker_color=status_colors[status],
                )
            )
        category_chart.update_layout(
            barmode="stack",
            height=350,
            xaxis_title="Resume category",
            yaxis_title="Candidates",
            legend_title="Decision",
        )
        st.plotly_chart(
            add_chart_hover(category_chart),
            alt="Candidate outcome counts grouped by resume category",
        )

    with rate_col:
        rate_chart = go.Figure(
            go.Bar(
                x=category_names,
                y=category_summary["Shortlist_Rate"].tolist(),
                marker_color="#337f88",
                text=[
                    f"{rate:.1f}%"
                    for rate in category_summary["Shortlist_Rate"]
                ],
                textposition="outside",
            )
        )
        rate_chart.update_layout(
            height=350,
            xaxis_title="Resume category",
            yaxis_title="Shortlist rate (%)",
            yaxis_range=[0, 100],
            showlegend=False,
        )
        st.plotly_chart(
            add_chart_hover(rate_chart),
            alt="Shortlist rate by resume category",
        )

    st.subheader("Most matched and missing skills")
    matched_counts = count_skills_by_candidate(results_df["Matched_Skills"])
    missing_counts = count_skills_by_candidate(results_df["Missing_Skills"])
    matched_col, missing_col = st.columns(2)

    with matched_col:
        st.markdown("**Most commonly matched**")
        if matched_counts:
            skills, counts = zip(*reversed(matched_counts))
            matched_chart = go.Figure(
                go.Bar(
                    x=list(counts),
                    y=list(skills),
                    orientation="h",
                    marker_color="#2e8b57",
                )
            )
            matched_chart.update_layout(
                height=360,
                xaxis_title="Resumes",
                yaxis_title="",
            )
            st.plotly_chart(
                add_chart_hover(matched_chart),
                alt="Most commonly matched job skills",
            )
        else:
            st.info("No matched skills were recorded.")

    with missing_col:
        st.markdown("**Most commonly missing**")
        if missing_counts:
            skills, counts = zip(*reversed(missing_counts))
            missing_chart = go.Figure(
                go.Bar(
                    x=list(counts),
                    y=list(skills),
                    orientation="h",
                    marker_color="#c85151",
                )
            )
            missing_chart.update_layout(
                height=360,
                xaxis_title="Resumes",
                yaxis_title="",
            )
            st.plotly_chart(
                add_chart_hover(missing_chart),
                alt="Most commonly missing job skills",
            )
        else:
            st.info("No missing skills were recorded.")

    st.subheader("Top-ranked candidates")
    top_candidates = results_df.nlargest(10, "Final_Score").sort_values(
        "Final_Score",
        ascending=True,
    )
    ranking_chart = go.Figure()
    candidate_labels = [
        f"{row.Rank}. {row.Resume}"
        for row in top_candidates.itertuples()
    ]
    for label, column in [
        ("Skills", "Skill_Match_%"),
        ("Experience", "Experience_Match_%"),
        ("Education", "Education_Match_%"),
        ("Final score", "Final_Score"),
    ]:
        ranking_chart.add_trace(
            go.Bar(
                name=label,
                y=candidate_labels,
                x=top_candidates[column].tolist(),
                orientation="h",
            )
        )
    ranking_chart.update_layout(
        barmode="group",
        height=max(360, len(top_candidates) * 42),
        xaxis_title="Score (%)",
        yaxis_title="",
        xaxis_range=[0, 100],
        legend_title="Match",
    )
    st.plotly_chart(
        add_chart_hover(ranking_chart),
        alt="Top candidates compared by skill, experience, education, and final scores",
    )

